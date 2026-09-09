"""
One-dimensional acoustic FDTD simulation for the experimental layer stack.

The geometry and material values are read from ``US_MRI_GROUP.xlsx``:

    water -> soft tissue -> water -> bone

Unlike the original two-medium example, this program advances acoustic pressure
``p`` and particle velocity ``v`` on a staggered grid:

    dv/dt = -(1/rho) dp/dx
    dp/dt = -K dv/dx

where rho = Z/c and K = rho*c**2 = Z*c.  Using both density and bulk modulus
means that interface reflections are set by acoustic impedance, rather than by
sound speed alone.  For pressure, the normal-incidence reflection coefficient
from material 1 to material 2 is (Z2 - Z1)/(Z2 + Z1).

Examples
--------
    python 1DWaveSimulation_Multilayer.py --animate
    python 1DWaveSimulation_Multilayer.py --animate --interval 15
    python 1DWaveSimulation_Multilayer.py

The pulse is a Gaussian-windowed sinusoid centered near 2 MHz, consistent with
the initial pulse in the workbook's ``Oscilloscope results`` sheet.  Pressure is
normalized because the spreadsheet voltage is a receiver response, not a
calibrated acoustic pressure.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.animation import FuncAnimation


DEFAULT_WORKBOOK = Path(__file__).with_name("US_MRI_GROUP.xlsx")


@dataclass(frozen=True)
class Material:
    name: str
    speed: float       # m/s
    impedance: float   # Pa*s/m (Rayl)
    color: str

    @property
    def density(self) -> float:
        return self.impedance / self.speed

    @property
    def bulk_modulus(self) -> float:
        return self.impedance * self.speed


@dataclass(frozen=True)
class ModelInputs:
    water_before_tissue: float
    soft_tissue_depth: float
    water_before_bone: float
    bone_depth: float
    water: Material
    soft_tissue: Material
    bone: Material
    stated_r_ws: float | None
    stated_r_wb: float | None

    @property
    def interfaces(self) -> tuple[float, float, float]:
        first = self.water_before_tissue
        second = first + self.soft_tissue_depth
        third = second + self.water_before_bone
        return first, second, third

    @property
    def length(self) -> float:
        return sum(
            (
                self.water_before_tissue,
                self.soft_tissue_depth,
                self.water_before_bone,
                self.bone_depth,
            )
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="1D acoustic simulation of water, soft tissue, water, and bone"
    )
    parser.add_argument(
        "--animate",
        action="store_true",
        help="Animate propagation and reflections instead of plotting snapshots",
    )
    parser.add_argument(
        "--interval",
        type=int,
        default=20,
        help="Delay between animation frames in milliseconds (default: 20)",
    )
    parser.add_argument(
        "--frame-stride",
        type=int,
        default=10,
        help="FDTD time steps calculated per animation frame (default: 10)",
    )
    parser.add_argument(
        "--duration-us",
        type=float,
        default=180.0,
        help="Simulated duration in microseconds (default: 180)",
    )
    parser.add_argument(
        "--dx-mm",
        type=float,
        default=0.05,
        help="Requested spatial grid spacing in millimeters (default: 0.05)",
    )
    parser.add_argument(
        "--frequency-mhz",
        type=float,
        default=2.0,
        help="Pulse center frequency in MHz (default: 2.0)",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Workbook containing the geometry and acoustic values",
    )
    return parser.parse_args()


def _sheet_case_insensitive(workbook, requested_name: str):
    matches = {
        sheet_name.strip().casefold(): sheet_name for sheet_name in workbook.sheetnames
    }
    actual_name = matches.get(requested_name.strip().casefold())
    if actual_name is None:
        raise ValueError(
            f"Workbook sheet {requested_name!r} was not found. "
            f"Available sheets: {', '.join(workbook.sheetnames)}"
        )
    return workbook[actual_name]


def _two_column_values(sheet) -> dict[str, float]:
    values: dict[str, float] = {}
    for key, value, *_ in sheet.iter_rows(values_only=True):
        if key is not None and value is not None:
            values[str(key).strip()] = float(value)
    return values


def load_model_inputs(workbook_path: Path) -> ModelInputs:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit(
            "Reading the experiment workbook requires openpyxl. "
            "Install it with: python -m pip install openpyxl"
        ) from exc

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path.resolve()}")

    workbook = openpyxl.load_workbook(
        workbook_path, read_only=True, data_only=True
    )
    try:
        distances = _two_column_values(
            _sheet_case_insensitive(workbook, "Distances between mediums")
        )
        values = _two_column_values(_sheet_case_insensitive(workbook, "Other values"))
    finally:
        workbook.close()

    required_distances = ("A", "B", "C", "D")
    required_values = ("c_w", "c_s", "c_b", "Z_w", "Z_s", "Z_b")
    missing = [key for key in required_distances if key not in distances]
    missing += [key for key in required_values if key not in values]
    if missing:
        raise ValueError(f"Missing workbook input(s): {', '.join(missing)}")

    if any(distances[key] <= 0.0 for key in required_distances):
        raise ValueError("All four layer distances must be positive")
    if any(values[key] <= 0.0 for key in required_values):
        raise ValueError("All sound speeds and acoustic impedances must be positive")

    return ModelInputs(
        water_before_tissue=distances["A"],
        soft_tissue_depth=distances["B"],
        water_before_bone=distances["C"],
        bone_depth=distances["D"],
        water=Material("Water", values["c_w"], values["Z_w"], "#d9f0ff"),
        soft_tissue=Material(
            "Soft tissue", values["c_s"], values["Z_s"], "#f7d6dc"
        ),
        bone=Material("Bone", values["c_b"], values["Z_b"], "#e8dfc7"),
        stated_r_ws=values.get("R_ws"),
        stated_r_wb=values.get("R_wb"),
    )


def pressure_reflection(z_from: float, z_to: float) -> float:
    return (z_to - z_from) / (z_to + z_from)


def print_model_summary(model: ModelInputs, dx: float, dt: float, steps: int) -> None:
    b1, b2, b3 = model.interfaces
    print("Layer stack loaded from workbook:")
    print(f"  Water:       0.00 to {b1 * 1e3:.2f} mm")
    print(f"  Soft tissue: {b1 * 1e3:.2f} to {b2 * 1e3:.2f} mm")
    print(f"  Water:       {b2 * 1e3:.2f} to {b3 * 1e3:.2f} mm")
    print(f"  Bone:        {b3 * 1e3:.2f} to {model.length * 1e3:.2f} mm")

    r_ws = pressure_reflection(model.water.impedance, model.soft_tissue.impedance)
    r_sw = pressure_reflection(model.soft_tissue.impedance, model.water.impedance)
    r_wb = pressure_reflection(model.water.impedance, model.bone.impedance)
    print("Pressure reflection coefficients calculated from Z:")
    print(f"  water -> soft tissue: {r_ws:+.4f}")
    print(f"  soft tissue -> water: {r_sw:+.4f}")
    print(f"  water -> bone:        {r_wb:+.4f}")
    if model.stated_r_ws is not None and model.stated_r_wb is not None:
        print(
            "  workbook rounded values: "
            f"R_ws={model.stated_r_ws:.3f}, R_wb={model.stated_r_wb:.3f}"
        )
    print(
        f"Grid: dx={dx * 1e3:.4f} mm, dt={dt * 1e9:.3f} ns, "
        f"steps={steps}"
    )


def material_arrays(model: ModelInputs, x: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    b1, b2, b3 = model.interfaces
    material_index = np.select(
        (x < b1, x < b2, x < b3),
        (0, 1, 0),
        default=2,
    )
    materials = (model.water, model.soft_tissue, model.bone)
    density = np.array([materials[i].density for i in material_index])
    bulk_modulus = np.array([materials[i].bulk_modulus for i in material_index])
    return density, bulk_modulus


def make_grid(
    model: ModelInputs, requested_dx: float, duration: float
) -> tuple[np.ndarray, np.ndarray, np.ndarray, float, float, int]:
    if requested_dx <= 0.0 or duration <= 0.0:
        raise ValueError("Grid spacing and duration must be positive")

    point_count = max(5, int(np.ceil(model.length / requested_dx)) + 1)
    x_pressure = np.linspace(0.0, model.length, point_count)
    dx = float(x_pressure[1] - x_pressure[0])
    x_velocity = 0.5 * (x_pressure[:-1] + x_pressure[1:])

    _, bulk_modulus = material_arrays(model, x_pressure)
    density_velocity, _ = material_arrays(model, x_velocity)

    max_speed = max(model.water.speed, model.soft_tissue.speed, model.bone.speed)
    dt_limit = dx / max_speed
    steps = int(np.ceil(duration / (0.95 * dt_limit)))
    dt = duration / steps
    if max_speed * dt / dx >= 1.0:
        raise RuntimeError("CFL stability condition was not satisfied")

    return x_pressure, density_velocity, bulk_modulus, dx, dt, steps


def pulse_profile(x: np.ndarray, center: float, sigma: float, wavenumber: float) -> np.ndarray:
    envelope = np.exp(-0.5 * ((x - center) / sigma) ** 2)
    return envelope * np.cos(wavenumber * (x - center))


def initial_fields(
    model: ModelInputs,
    x_pressure: np.ndarray,
    dx: float,
    dt: float,
    frequency: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wavelength = model.water.speed / frequency
    if dx > wavelength / 10.0:
        raise ValueError(
            f"The {dx * 1e3:.4f} mm grid is too coarse for "
            f"{frequency / 1e6:.3f} MHz. Use --dx-mm {wavelength * 100:.4f} or smaller."
        )

    # Keep the pulse well inside the first water layer.  Its Gaussian sigma is
    # 1.25 wavelengths, giving a short experimental-style tone burst.
    sigma = 1.25 * wavelength
    center = min(5.0e-3, 0.30 * model.water_before_tissue)
    if center < 3.5 * sigma:
        center = 3.5 * sigma
    if center + 3.5 * sigma >= model.water_before_tissue:
        raise ValueError("The first water layer is too short to contain the source pulse")

    wavenumber = 2.0 * np.pi / wavelength
    pressure = pulse_profile(x_pressure, center, sigma, wavenumber)

    # Velocity lives half a cell and half a time step behind pressure.  Sampling
    # the analytic right-going pulse at t=-dt/2 suppresses a spurious left-going
    # component at startup.
    x_velocity = 0.5 * (x_pressure[:-1] + x_pressure[1:])
    velocity = pulse_profile(
        x_velocity + model.water.speed * dt / 2.0,
        center,
        sigma,
        wavenumber,
    ) / model.water.impedance
    pressure_next = np.empty_like(pressure)
    return pressure, velocity, pressure_next


def advance_acoustics(
    pressure: np.ndarray,
    velocity: np.ndarray,
    pressure_next: np.ndarray,
    density_velocity: np.ndarray,
    bulk_modulus: np.ndarray,
    dt: float,
    dx: float,
    left_speed: float,
    right_speed: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Advance the staggered pressure/velocity fields by one stable time step."""
    velocity -= (dt / (density_velocity * dx)) * np.diff(pressure)
    pressure_next[1:-1] = pressure[1:-1] - (
        bulk_modulus[1:-1] * dt / dx
    ) * np.diff(velocity)

    # First-order Mur conditions approximate an open continuation of the first
    # water layer on the left and the final bone layer on the right.
    alpha_left = (left_speed * dt - dx) / (left_speed * dt + dx)
    alpha_right = (right_speed * dt - dx) / (right_speed * dt + dx)
    pressure_next[0] = pressure[1] + alpha_left * (
        pressure_next[1] - pressure[0]
    )
    pressure_next[-1] = pressure[-2] + alpha_right * (
        pressure_next[-2] - pressure[-1]
    )
    return pressure_next, pressure


def decorate_axes(ax, model: ModelInputs) -> None:
    b1, b2, b3 = model.interfaces
    edges = (0.0, b1, b2, b3, model.length)
    layers = (model.water, model.soft_tissue, model.water, model.bone)
    labels = ("Water", "Soft tissue", "Water", "Bone plate")

    for left, right, material, label in zip(edges[:-1], edges[1:], layers, labels):
        ax.axvspan(left * 1e3, right * 1e3, color=material.color, alpha=0.55)
        ax.text(
            (left + right) * 0.5e3,
            1.47,
            label,
            ha="center",
            va="center",
            fontsize=9,
        )
    for boundary in (b1, b2, b3):
        ax.axvline(boundary * 1e3, color="black", linestyle="--", alpha=0.55)

    ax.set_xlim(0.0, model.length * 1e3)
    ax.set_ylim(-1.65, 1.65)
    ax.set_xlabel("Distance from transducer (mm)")
    ax.set_ylabel("Normalized acoustic pressure")
    ax.grid(True, alpha=0.25)


def arrival_times(model: ModelInputs, source_center: float = 5.0e-3) -> dict[str, float]:
    b1, b2, b3 = model.interfaces
    t_soft = max(0.0, b1 - source_center) / model.water.speed
    t_water_2 = t_soft + (b2 - b1) / model.soft_tissue.speed
    t_bone = t_water_2 + (b3 - b2) / model.water.speed
    return {
        "Initial pulse": 0.0,
        "At soft tissue": t_soft + 4.0e-6,
        "Back in water": t_water_2 + 4.0e-6,
        "At bone": t_bone + 4.0e-6,
        "Bone echo returning": 2.0 * t_bone,
    }


def run_snapshots(
    model: ModelInputs,
    pressure: np.ndarray,
    velocity: np.ndarray,
    pressure_next: np.ndarray,
    density_velocity: np.ndarray,
    bulk_modulus: np.ndarray,
    dt: float,
    dx: float,
    steps: int,
) -> dict[str, tuple[float, np.ndarray]]:
    requested = arrival_times(model)
    snapshots: dict[str, tuple[float, np.ndarray]] = {
        "Initial pulse": (0.0, pressure.copy())
    }
    pending = [
        (label, time_value)
        for label, time_value in requested.items()
        if label != "Initial pulse" and time_value <= steps * dt
    ]
    pending.sort(key=lambda item: item[1])

    for step in range(1, steps + 1):
        pressure, pressure_next = advance_acoustics(
            pressure,
            velocity,
            pressure_next,
            density_velocity,
            bulk_modulus,
            dt,
            dx,
            model.water.speed,
            model.bone.speed,
        )
        while pending and step * dt >= pending[0][1]:
            label, _ = pending.pop(0)
            snapshots[label] = (step * dt, pressure.copy())

    return snapshots


def plot_snapshots(
    x_pressure: np.ndarray,
    model: ModelInputs,
    snapshots: dict[str, tuple[float, np.ndarray]],
) -> None:
    fig, ax = plt.subplots(figsize=(12, 6.5))
    decorate_axes(ax, model)
    colors = plt.cm.viridis(np.linspace(0.05, 0.9, len(snapshots)))
    for color, (label, (time_value, pressure)) in zip(colors, snapshots.items()):
        ax.plot(
            x_pressure * 1e3,
            pressure,
            color=color,
            lw=1.7,
            label=f"{label} ({time_value * 1e6:.1f} us)",
        )
    ax.set_title("Acoustic propagation through water, soft tissue, water, and bone")
    ax.legend(loc="lower left", fontsize=8)
    fig.tight_layout()
    plt.show()


def animate_simulation(
    x_pressure: np.ndarray,
    model: ModelInputs,
    pressure: np.ndarray,
    velocity: np.ndarray,
    pressure_next: np.ndarray,
    density_velocity: np.ndarray,
    bulk_modulus: np.ndarray,
    dt: float,
    dx: float,
    steps: int,
    interval: int,
    frame_stride: int,
) -> None:
    if interval <= 0 or frame_stride <= 0:
        raise ValueError("Animation interval and frame stride must be positive")

    fig, ax = plt.subplots(figsize=(12, 6.5))
    decorate_axes(ax, model)
    (line,) = ax.plot(x_pressure * 1e3, pressure, color="#1261a0", lw=1.8)
    time_text = ax.text(0.015, 0.04, "", transform=ax.transAxes)
    ax.set_title("Acoustic propagation and impedance reflections")
    fig.tight_layout()

    completed_steps = 0

    def update(frame):
        nonlocal pressure, pressure_next, completed_steps
        # Advance to an absolute target so repeated draws of the same animation
        # frame (which some Matplotlib backends request) cannot advance time twice.
        target_steps = min(frame * frame_stride, steps)
        for _ in range(target_steps - completed_steps):
            pressure, pressure_next = advance_acoustics(
                pressure,
                velocity,
                pressure_next,
                density_velocity,
                bulk_modulus,
                dt,
                dx,
                model.water.speed,
                model.bone.speed,
            )
            completed_steps += 1
        line.set_ydata(pressure)
        time_text.set_text(f"Time: {completed_steps * dt * 1e6:7.2f} us")
        return line, time_text

    frame_count = int(np.ceil(steps / frame_stride)) + 1
    animation = FuncAnimation(
        fig,
        update,
        frames=frame_count,
        interval=interval,
        repeat=False,
        blit=False,
    )
    fig._animation = animation
    plt.show()


def main() -> None:
    args = parse_args()
    if args.frequency_mhz <= 0.0:
        raise ValueError("Pulse frequency must be positive")

    model = load_model_inputs(args.workbook)
    duration = args.duration_us * 1e-6
    requested_dx = args.dx_mm * 1e-3
    frequency = args.frequency_mhz * 1e6
    (
        x_pressure,
        density_velocity,
        bulk_modulus,
        dx,
        dt,
        steps,
    ) = make_grid(model, requested_dx, duration)
    pressure, velocity, pressure_next = initial_fields(
        model, x_pressure, dx, dt, frequency
    )
    print_model_summary(model, dx, dt, steps)

    if args.animate:
        animate_simulation(
            x_pressure,
            model,
            pressure,
            velocity,
            pressure_next,
            density_velocity,
            bulk_modulus,
            dt,
            dx,
            steps,
            args.interval,
            args.frame_stride,
        )
    else:
        snapshots = run_snapshots(
            model,
            pressure,
            velocity,
            pressure_next,
            density_velocity,
            bulk_modulus,
            dt,
            dx,
            steps,
        )
        plot_snapshots(x_pressure, model, snapshots)


if __name__ == "__main__":
    main()
